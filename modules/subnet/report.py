"""Word-report export for reviewed VLSM calculation results."""
from pathlib import Path
from docx import Document


def export_vlsm_docx(destination, base_network, allocations):
    """Create a local Word report; callers must provide already-calculated data."""
    path = Path(destination)
    if path.suffix.lower() != '.docx':
        raise ValueError('Choose a .docx report file.')
    report = Document()
    report.add_heading('Bastion — VLSM Allocation Report', 0)
    report.add_paragraph(f'Base network: {base_network}')
    table = report.add_table(rows=1, cols=6)
    for cell, label in zip(table.rows[0].cells, ('Requirement', 'Requested hosts', 'Network', 'CIDR', 'Usable range', 'Broadcast')):
        cell.text = label
    for item in allocations:
        cells = table.add_row().cells
        values = (item['name'], str(item['requested_hosts']), item['network'], item['cidr'], item['usable_range'], item['broadcast'])
        for cell, value in zip(cells, values): cell.text = value
    report.save(path)
    return path
